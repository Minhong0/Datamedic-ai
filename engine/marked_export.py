"""상태 표시가 들어간 엑셀 내보내기."""
from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from engine.json_io import apply_cleaned_to_json, is_json_file, row_label
from engine.masking import BIZNO_COLUMN, COMPANY_COLUMN, mask_cell_value
from engine.models import Dataset
from engine.viewer_marks import CHECK_LABELS, CellMark, MarkDetail, ViewerMarks, rule_label

AUTHOR = "DataMedic AI"
EXTRA_SHEETS = ("범례", "검사 요약", "삭제된 행", "변경 이력", "회신 대기")
STATUS_LABELS = {
    "auto": "자동 수정", "auto_pending": "자동 수정 예정", "approval": "승인 대기", "undecided": "미결정",
    "operator": "담당자 수정", "rejected": "거절", "request": "확인 요청",
}
_WHO = {"system": "시스템", "operator": "담당자"}
_INT_RE = re.compile(r"^-?\d+$")
AMOUNT_COLUMN = "지원금액"

FILLS = {
    "auto": "EEF8F1", "approval": "FFF6D6", "operator": "E3EEFB", "request": "FDF0DC", "undecided": "F0F0F0",
}
FONT_COLORS = {"approval": "7A5B00", "operator": "1F4E9A"}
GROUP_BORDER = "6E56CF"
SENSITIVE = (BIZNO_COLUMN, COMPANY_COLUMN)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def _value(column: str, raw: object, *, masked: bool, amount_cols: set[str]):
    """셀에 쓸 값 — 금액 열의 정수 문자열은 숫자 셀로, 발표용이면 사업자번호·기업명 열은 마스킹."""
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return None
    text = str(raw)
    if masked and column in SENSITIVE:
        return mask_cell_value(column, text)
    if column in amount_cols and _INT_RE.match(text.strip()):
        return int(text.strip())
    return text


def _memo(mark: CellMark | None, group_labels: list[str]) -> str:
    parts: list[str] = []
    for d in (mark.details if mark else []):
        line = d.reason
        if d.decided_by:
            line += f" (처리자: {_WHO.get(d.decided_by, d.decided_by)})"
        parts.append(line)
    parts += [f"묶음 규칙 · {label}" for label in group_labels]
    return "\n".join(parts)


def _comment(text: str) -> Comment:
    c = Comment(text, AUTHOR)
    c.width, c.height = 320, max(60, 22 * (text.count("\n") + 2))
    return c


def _style_cell(cell, mark: CellMark) -> None:
    status = mark.status
    if status in FILLS:
        cell.fill = _fill(FILLS[status])
    color = FONT_COLORS.get(status)
    edited = status == "operator" and any(d.edited for d in mark.details)
    if color or edited:
        cell.font = Font(color=color, bold=edited)


def _auto_width(ws, header_len: dict[int, int]) -> None:
    for col_idx, rows in header_len.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(rows + 2, 8), 48)


def _write_data_sheet(wb, file: str, sheet: str, original: pd.DataFrame, cleaned: pd.DataFrame | None,
                      marks: ViewerMarks, *, masked: bool, amount_cols: set[str]) -> None:
    ws = wb.create_sheet(sheet)
    cols = [c for c in original.columns if c != "_row"]
    ws.append(cols)
    for i in range(1, len(cols) + 1):
        ws.cell(row=1, column=i).font = Font(bold=True)
    ws.freeze_panes = "A2"
    widths = {i + 1: len(str(c)) for i, c in enumerate(cols)}

    source = cleaned if cleaned is not None else original
    purple = Side(style="medium", color=GROUP_BORDER)
    for _, row in source.iterrows():
        orig_row = int(row["_row"])
        ws.append([_value(c, row[c] if c in row.index else None, masked=masked, amount_cols=amount_cols) for c in cols])
        r = ws.max_row
        group = marks.group_rows.get((file, sheet, orig_row), [])
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=i)
            widths[i] = max(widths[i], len(str(cell.value)) if cell.value is not None else 0)
            mark = marks.cells.get((file, sheet, orig_row, c))
            if mark is not None:
                _style_cell(cell, mark)
            memo = _memo(mark, group if i == 1 else [])
            if memo:
                cell.comment = _comment(memo)
            if group and i == 1:
                cell.border = Border(left=purple, right=purple, top=purple, bottom=purple)
    _auto_width(ws, widths)


def _write_legend(wb, file: str, *, masked: bool) -> None:
    ws = wb.create_sheet("범례", 0)
    ws.append([f"{file} — 변경 내역 표시 범례"])
    ws["A1"].font = Font(bold=True, size=13)
    if masked:
        ws.append(["발표용 — 마스킹됨 (사업자번호·기업명 열의 값이 가려져 있습니다. 업무 데이터로 쓰지 마세요.)"])
        ws["A2"].font = Font(bold=True, color="B00020")
    ws.append([f"생성: {datetime.now():%Y-%m-%d %H:%M} · 상태 설명은 셀 메모(빨간 삼각형)에 있습니다."])
    ws.append([])
    ws.append(["표시", "의미"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    rows = [
        ("auto", "자동 수정 — 메모에 이전 값·규칙"),
        ("approval", "승인 대기 — 메모에 제안값·사유 (확정 전 내보낸 경우에만)"),
        ("operator", "담당자 승인·직접 수정 — 직접 입력은 굵게, 메모에 처리자·이전 값"),
        ("request", "확인 요청 — 회신 대기 (회신 대기 시트에 목록)"),
        ("undecided", "미결정 — 결정 없이 끝나 원래 값 유지"),
    ]
    for status, text in rows:
        ws.append(["샘플", text])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.fill = _fill(FILLS[status])
        if status in FONT_COLORS:
            cell.font = Font(color=FONT_COLORS[status], bold=(status == "operator"))
    ws.append(["샘플", "거절 — 원래 값 유지, 거절된 제안은 메모에만 (배경 없음)"])
    ws.append(["샘플", "묶음 규칙 대상 행 — 첫 열 보라 테두리, 메모에 묶음 라벨"])
    purple = Side(style="medium", color=GROUP_BORDER)
    ws.cell(row=ws.max_row, column=1).border = Border(left=purple, right=purple, top=purple, bottom=purple)
    ws.append(["", "삭제된 중복 행은 정제본에서 빠지고 '삭제된 행' 시트에 기록됩니다."])
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 80


def _rows_for_file(marks: ViewerMarks, file: str):
    for (f, s, r, c), mark in sorted(marks.cells.items(), key=lambda kv: (kv[0][1], kv[0][2], kv[0][3])):
        if f == file:
            yield s, r, c, mark


def _write_extra_sheets(wb, file: str, original_by_sheet: dict[str, pd.DataFrame], marks: ViewerMarks,
                        *, masked: bool) -> None:
    ws = wb.create_sheet("검사 요약")
    ws.append(["시트", "검사", "결과", "건수", "확인 요청", "문제 없음", "비고"])
    any_check = False
    for (f, s), checks in sorted(marks.sheet_checks.items()):
        if f != file:
            continue
        for c in checks:
            any_check = True
            result = "통과 ✓" if c.issue_count == 0 else "확인 필요"
            ws.append([s, c.label, result, c.issue_count, c.status_counts.get("request", 0),
                       c.status_counts.get("ok", 0), "수정 후 다시 집계 필요" if c.stale else ""])
    if not any_check:
        ws.append(["", "(집계 검사를 실행하지 않았습니다)", "", "", "", "", ""])

    ws = wb.create_sheet("삭제된 행")
    ws.append(["시트", "원본 행", "원래 내용", "사유", "처리자"])
    for (f, s, r), d in sorted(marks.deleted_rows.items(), key=lambda kv: (kv[0][1], kv[0][2])):
        if f != file:
            continue
        df = original_by_sheet.get(s)
        content = ""
        if df is not None:
            hit = df[df["_row"].astype(int) == r]
            if not hit.empty:
                content = " | ".join(
                    str(mask_cell_value(c, str(hit.iloc[0][c])) or "") for c in df.columns if c != "_row")
        ws.append([s, r, content, d.reason, _WHO.get(d.decided_by or "", d.decided_by or "")])

    ws = wb.create_sheet("변경 이력")
    ws.append(["시트", "원본 행", "열", "상태", "규칙", "이전", "이후/제안", "사유", "처리자", "시각"])
    for s, r, c, mark in _rows_for_file(marks, file):
        for d in mark.details:
            ws.append([s, r, c, STATUS_LABELS[d.status], rule_label(d.rule_id), d.before or "", d.after or "", d.reason,
                       _WHO.get(d.decided_by or "", d.decided_by or ""),
                       d.timestamp.strftime("%Y-%m-%d %H:%M:%S") if d.timestamp else ""])

    ws = wb.create_sheet("회신 대기")
    ws.append(["시트", "원본 행", "열", "현재 값", "요청 사유", "요청 기관(파일)"])
    org = Path(file).stem
    for s, r, c, mark in _rows_for_file(marks, file):
        for d in mark.details:
            if d.status == "request":
                ws.append([s, r, c, d.before or "", d.reason, org])
    for name in EXTRA_SHEETS[1:]:
        sheet = wb[name]
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for i in range(1, sheet.max_column + 1):
            longest = max((len(str(sheet.cell(row=j, column=i).value or "")) for j in range(1, sheet.max_row + 1)), default=8)
            sheet.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 8), 60)


def build_marked_workbooks(
    original: Dataset,
    cleaned: Dataset | None,
    marks: ViewerMarks,
    *,
    masked: bool = False,
    amount_columns: dict[tuple[str, str], set[str]] | None = None,
) -> dict[str, bytes]:
    """파일마다 상태 표시가 들어간 엑셀을 만든다. 반환: {파일 이름: xlsx 바이트} (발표용이면 이름에 `_masked`)."""
    amount_columns = amount_columns or {}
    out: dict[str, bytes] = {}
    for file in sorted({f for f, _ in original.tables if not is_json_file(f)}):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        sheets = sorted(s for f, s in original.tables if f == file)
        for sheet in sheets:
            amt = set(amount_columns.get((file, sheet), {AMOUNT_COLUMN}))
            _write_data_sheet(wb, file, sheet, original.tables[(file, sheet)],
                              cleaned.tables.get((file, sheet)) if cleaned is not None else None, marks,
                              masked=masked, amount_cols=amt)
        _write_extra_sheets(wb, file, {s: original.tables[(file, s)] for s in sheets}, marks, masked=masked)
        _write_legend(wb, file, masked=masked)
        buf = io.BytesIO()
        wb.save(buf)
        name = Path(file)
        out[f"{name.stem}_masked{name.suffix}" if masked else file] = buf.getvalue()
    return out


def save_marked_workbooks(files: dict[str, bytes], run_dir: Path, *, masked: bool) -> list[Path]:
    """만들어진 엑셀을 outputs/<run_id>/marked/ (발표용은 marked_masked/)에 저장한다."""
    out_dir = run_dir / ("marked_masked" if masked else "marked")
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, data in files.items():
        p = out_dir / name
        p.write_bytes(data)
        paths.append(p)
    return paths


def changes_csv(marks: ViewerMarks, file: str) -> bytes:
    """변경 이력 CSV (엑셀에서 바로 열리도록 UTF-8 BOM). 위치는 JSON 이면 배열 인덱스 `[0]`. 값은 이미 마스킹되어 있다."""
    import csv

    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["시트", "행", "열", "상태", "규칙", "이전", "이후/제안", "사유", "처리자", "시각"])
    for (f, s, r, c), mark in sorted(marks.cells.items(), key=lambda kv: (kv[0][1], kv[0][2], kv[0][3])):
        if f != file:
            continue
        for d in mark.details:
            w.writerow([s, row_label(f, r), c, STATUS_LABELS[d.status], rule_label(d.rule_id), d.before or "", d.after or "",
                        d.reason, _WHO.get(d.decided_by or "", d.decided_by or ""),
                        d.timestamp.strftime("%Y-%m-%d %H:%M:%S") if d.timestamp else ""])
    for (f, s, r), d in sorted(marks.deleted_rows.items(), key=lambda kv: (kv[0][1], kv[0][2])):
        if f == file:
            w.writerow([s, row_label(f, r), "", "삭제된 행", rule_label(d.rule_id), d.before or "", "", d.reason,
                        _WHO.get(d.decided_by or "", d.decided_by or ""), ""])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def build_json_exports(
    original: Dataset,
    cleaned: Dataset | None,
    marks: ViewerMarks,
    input_paths: list[str],
    *,
    masked: bool = False,
    amount_columns: dict[tuple[str, str], set[str]] | None = None,
) -> dict[str, bytes]:
    """JSON 입력 파일마다 정제된 JSON(원래 형식 그대로)과 변경 이력 CSV 를 만든다 (발표용이면 `_masked`)."""
    out: dict[str, bytes] = {}
    paths = {Path(p).name: p for p in input_paths}
    for file in sorted({f for f, _ in original.tables if is_json_file(f)}):
        src = paths.get(file)
        if src is None:
            continue
        sheets = sorted(sh for f, sh in original.tables if f == file)
        amt: set[str] = set()
        for sh in sheets:
            amt |= (amount_columns or {}).get((file, sh), set())
        data = apply_cleaned_to_json(
            src, {sh: original.tables[(file, sh)] for sh in sheets},
            {sh: cleaned.tables[(file, sh)] for sh in sheets if cleaned is not None and (file, sh) in cleaned.tables},
            amount_columns=amt or None, masked=masked)
        stem = Path(file).stem + ("_masked" if masked else "")
        out[f"{stem}.json"] = data
        out[f"{stem}_변경이력.csv"] = changes_csv(marks, file)
    return out


def build_export_bundle(
    original: Dataset,
    cleaned: Dataset | None,
    marks: ViewerMarks,
    *,
    input_paths: list[str] | None = None,
    masked: bool = False,
    amount_columns: dict[tuple[str, str], set[str]] | None = None,
) -> dict[str, bytes]:
    """내려받을 파일 전체 — 엑셀은 상태 표시 엑셀, JSON 은 원래 형식의 JSON + 변경 이력 CSV."""
    files = build_marked_workbooks(original, cleaned, marks, masked=masked, amount_columns=amount_columns)
    files.update(build_json_exports(original, cleaned, marks, input_paths or [], masked=masked,
                                    amount_columns=amount_columns))
    return files


def zip_workbooks(files: dict[str, bytes]) -> bytes:
    """내려받기용 ZIP — 파일이 여럿일 때 한 번에 받는다."""
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    return buf.getvalue()


def amount_columns_from_profiles(profiles) -> dict[tuple[str, str], set[str]]:
    """구조 분석 결과(SheetProfile)에서 금액으로 분류된 열을 {(파일, 시트): 열 이름들}로. 없으면 기본 '지원금액'."""
    out: dict[tuple[str, str], set[str]] = {}
    for p in profiles or []:
        cols = {c.name for c in p.columns if c.semantic_type == "amount"}
        file = getattr(p, "file", None)
        if cols and file is not None:
            out[(file, p.sheet)] = cols
    return out


__all__ = [
    "AUTHOR", "EXTRA_SHEETS", "FILLS", "build_marked_workbooks", "save_marked_workbooks",
    "amount_columns_from_profiles", "zip_workbooks", "build_json_exports", "build_export_bundle", "changes_csv", "CHECK_LABELS", "MarkDetail",
]
