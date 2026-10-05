"""깔끔하게 정리된 엑셀 내려받기 — 기관별 정제본 각각 + 전체를 합친 엑셀 하나."""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from engine.masking import BIZNO_COLUMN, COMPANY_COLUMN, mask_cell_value
from engine.viewer_marks import CHECK_LABELS

AMOUNT_COLUMN = "지원금액"
PENDING_TEXT = "(회신 대기)"
NOTE_COLUMN = "확인 요청"
_SCOPE_TEXT = {"per_business": "사업별 합계", "total": "모든 사업 합산", "per_payment": "건별"}
_AGG_SUFFIX = " (집계 확인)"
_PENDING_FILL = PatternFill("solid", fgColor="FFF2CC")
MERGED_NAME = "전체_합본_정제.xlsx"
MERGED_SHEET = "전체"
_HEADER_FILL = PatternFill("solid", fgColor="1F3A6E")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_MAX_WIDTH = 40


def _is_excel(path: Path) -> bool:
    return path.suffix.lower() == ".xlsx"


PendingCells = set[tuple[str, str, int, str]]
DeletedRows = set[tuple[str, str, int]]


def pending_cells(issues, requests, handled_orgs: set[str]) -> PendingCells:
    """보완 요청 메일을 이미 처리(발송 또는 시험 저장)한 오류 칸 — 내려받는 엑셀에서 '(회신 대기)'로 바꿀 곳."""
    by_id = {i.issue_id: i for i in issues if getattr(i, "ref", None) is not None}
    out: PendingCells = set()
    for req in requests:
        if "(집계" in req.org or not (req.sent or req.org in handled_orgs):
            continue
        for iid in req.issues:
            issue = by_id.get(iid)
            if issue is not None:
                ref = issue.ref
                out.add((Path(ref.file).name, ref.sheet, int(ref.row), ref.column))
    return out


def deleted_rows(changes) -> DeletedRows:
    """승인으로 삭제된 행 — 정제본에서는 사라져 있어서 원본 행 번호를 맞출 때 건너뛴다."""
    return {(Path(c.ref.file).name, c.ref.sheet, int(c.ref.row)) for c in changes if c.after is None}


def _with_pending(file: str, sheet: str, header: list[str], rows: list[list[object]],
                  pending: PendingCells | None, deleted: DeletedRows | None) -> list[list[object]]:
    """정제본의 몇 번째 줄이 원본의 몇 행인지 맞춰, 회신을 기다리는 칸을 '(회신 대기)'로 바꾼 새 목록."""
    mine = {(r, c) for f, s, r, c in (pending or ()) if f == file and s == sheet}
    if not mine:
        return rows
    gone = {r for f, s, r in (deleted or ()) if f == file and s == sheet}
    col_of = {h: i for i, h in enumerate(header)}
    out, orig = [], 2
    for row in rows:
        while orig in gone:
            orig += 1
        new = list(row)
        for col in [c for r, c in mine if r == orig and c in col_of]:
            new[col_of[col]] = PENDING_TEXT
        out.append(new)
        orig += 1
    return out


AggNotes = dict[tuple[str, str, int], list[str]]


def _short_org(file: str) -> str:
    return Path(file).stem.split("_")[0]


def aggregate_notes(agg_results: list[dict], agg_requests, handled_orgs: set[str]) -> AggNotes:
    """집계 확인 요청 메일을 처리한(발송 또는 시험 저장) 기관 파일의 행마다 '무슨 검사에 걸렸는지' 사유를 만든다."""
    groups = {f"{res['capability']}:{g.issue_id}": (res, g)
              for res in agg_results if not res.get("skipped") for g in res.get("group_issues", [])}
    notes: AggNotes = {}
    for req in agg_requests:
        if not req.org.endswith(_AGG_SUFFIX) or not (req.sent or req.org in handled_orgs):
            continue
        file = req.org[: -len(_AGG_SUFFIX)]
        for key in req.issues:
            found = groups.get(key)
            if found is None:
                continue
            res, g = found
            label = CHECK_LABELS.get(g.rule_id, g.rule_id)
            if g.rule_id == "SUM_LIMIT":
                scope = ((res.get("spec") and res["spec"].params) or {}).get("scope", "per_business")
                label += f" · {_SCOPE_TEXT.get(scope, scope)}"
            by_file: dict[str, list[int]] = {}
            for r in g.rows:
                by_file.setdefault(Path(r.file).name, []).append(int(r.row))
            where = " / ".join(f"{_short_org(f)} {'·'.join(str(n) for n in sorted(set(rows)))}행" for f, rows in by_file.items())
            text = f"{PENDING_TEXT} {label} — {where}"
            for r in g.rows:
                if Path(r.file).name == file:
                    slot = notes.setdefault((file, r.sheet, int(r.row)), [])
                    if text not in slot:
                        slot.append(text)
    return notes


def _with_notes(file: str, sheet: str, header: list[str], rows: list[list[object]],
                notes: AggNotes | None, deleted: DeletedRows | None) -> tuple[list[str], list[list[object]]]:
    """사유가 있는 파일·시트면 맨 끝에 '확인 요청' 열을 붙인다 (삭제된 행을 건너뛰며 원본 행 번호를 맞춘다)."""
    mine = {r: txt for (f, s, r), txt in (notes or {}).items() if f == file and s == sheet}
    if not mine or NOTE_COLUMN in header:
        return header, rows
    gone = {r for f, s, r in (deleted or ()) if f == file and s == sheet}
    out, orig = [], 2
    for row in rows:
        while orig in gone:
            orig += 1
        out.append([*row, "\n".join(mine[orig]) if orig in mine else None])
        orig += 1
    return [*header, NOTE_COLUMN], out


def _cell_value(column: str, raw: object, *, masked: bool):
    """화면용 값 — 지원금액은 숫자로, 발표용이면 사업자번호·기업명을 가린다."""
    if raw is None:
        return None
    if raw == PENDING_TEXT:
        return raw
    if masked and column in (BIZNO_COLUMN, COMPANY_COLUMN):
        return mask_cell_value(column, str(raw))
    if column == AMOUNT_COLUMN and isinstance(raw, str) and re.fullmatch(r"-?\d+", raw.strip()):
        return int(raw.strip())
    return raw


def _style_sheet(ws, *, amount_cols: set[int]) -> None:
    for cell in ws[1]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    for idx in range(1, ws.max_column + 1):
        letter = get_column_letter(idx)
        longest = max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(_MAX_WIDTH, max(10, longest + 3))
        if idx in amount_cols:
            for c in ws[letter][1:]:
                c.number_format = "#,##0"
                c.alignment = Alignment(horizontal="right")


def _read_sheets(path: Path) -> dict[str, tuple[list[str], list[list[object]]]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    out: dict[str, tuple[list[str], list[list[object]]]] = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h) if h is not None else "" for h in rows[0]]
        out[ws.title] = (header, [list(r) for r in rows[1:]])
    return out


def _write_sheet(wb: openpyxl.Workbook, title: str, header: list[str], rows: list[list[object]], *, masked: bool) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append([_cell_value(h, v, masked=masked) for h, v in zip(header, r)])
    _style_sheet(ws, amount_cols={i + 1 for i, h in enumerate(header) if h == AMOUNT_COLUMN or h == "지원금액합계"})
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.value == PENDING_TEXT:
                cell.fill, cell.alignment = _PENDING_FILL, Alignment(horizontal="center")
            elif isinstance(cell.value, str) and cell.value.startswith(PENDING_TEXT):
                cell.fill, cell.alignment = _PENDING_FILL, Alignment(wrap_text=True, vertical="top")
    if NOTE_COLUMN in header:
        ws.column_dimensions[get_column_letter(header.index(NOTE_COLUMN) + 1)].width = 52


def _to_bytes(wb: openpyxl.Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def tidy_name(file_name: str) -> str:
    """'거제시_사업실적_2026.xlsx' → '거제시_사업실적_2026_정제.xlsx'"""
    return f"{Path(file_name).stem}_정제.xlsx"


def build_clean_workbooks(cleaned_paths: list[Path], *, masked: bool = False,
                          pending: PendingCells | None = None, deleted: DeletedRows | None = None,
                          notes: AggNotes | None = None) -> dict[str, bytes]:
    """기관별 정제 엑셀 — {파일명: 엑셀 바이트}. 엑셀이 아닌 파일(JSON)은 건너뛴다."""
    files: dict[str, bytes] = {}
    for path in cleaned_paths:
        if not _is_excel(path):
            continue
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for title, (header, rows) in _read_sheets(path).items():
            rows = _with_pending(path.name, title, header, rows, pending, deleted)
            header, rows = _with_notes(path.name, title, header, rows, notes, deleted)
            _write_sheet(wb, title, header, rows, masked=masked)
        files[tidy_name(path.name)] = _to_bytes(wb)
    return files


def build_clean_files(cleaned_paths: list[Path], *, masked: bool = False,
                      pending: PendingCells | None = None, deleted: DeletedRows | None = None,
                      notes: AggNotes | None = None) -> dict[str, bytes]:
    """기관별 정제 파일 — 엑셀은 보기 좋게 다시 쓰고, JSON 입력의 정제본은 원래 JSON 그대로 담는다."""
    files = build_clean_workbooks(cleaned_paths, masked=masked, pending=pending, deleted=deleted, notes=notes)
    if not masked:
        for path in cleaned_paths:
            if path.suffix.lower() == ".json":
                files[f"{path.stem}_정제.json"] = path.read_bytes()
    return files


def build_merged_workbook(cleaned_paths: list[Path], *, masked: bool = False,
                          pending: PendingCells | None = None, deleted: DeletedRows | None = None,
                          notes: AggNotes | None = None) -> bytes | None:
    """모든 기관의 첫 번째 시트(실적)를 한 시트로 합친 엑셀 + 기관별 요약. 합칠 엑셀이 없으면 None."""
    header: list[str] = []
    merged: list[dict[str, object]] = []
    summary: list[list[object]] = []
    for path in cleaned_paths:
        if not _is_excel(path):
            continue
        sheets = _read_sheets(path)
        if not sheets:
            continue
        title, (head, rows) = next(iter(sheets.items()))
        rows = _with_pending(path.name, title, head, rows, pending, deleted)
        head, rows = _with_notes(path.name, title, head, rows, notes, deleted)
        for h in head:
            if h and h not in header:
                header.append(h)
        amount_total = 0
        for r in rows:
            rec = {h: v for h, v in zip(head, r)}
            merged.append(rec)
            val = _cell_value(AMOUNT_COLUMN, rec.get(AMOUNT_COLUMN), masked=False)
            amount_total += val if isinstance(val, int) else 0
        summary.append([path.stem, len(rows), amount_total])
    if not merged:
        return None
    if NOTE_COLUMN in header:
        header = [h for h in header if h != NOTE_COLUMN] + [NOTE_COLUMN]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, MERGED_SHEET, header, [[rec.get(h) for h in header] for rec in merged], masked=masked)
    _write_sheet(wb, "기관별 요약", ["파일", "행 수", "지원금액합계"], summary, masked=False)
    return _to_bytes(wb)


def zip_clean_workbooks(files: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    return buf.getvalue()
